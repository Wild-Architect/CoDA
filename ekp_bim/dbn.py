from __future__ import annotations

import json
import os
import subprocess
from pathlib import Path

from openpyxl import load_workbook


def load_catalog(path: Path) -> list[dict]:
    if not path.exists():
        return []
    if path.suffix.lower() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook["Каталог ДБН"] if "Каталог ДБН" in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows)]
        result = []
        for values in rows:
            row = dict(zip(headers, values))
            filename = str(row.get("Файл ДБН") or "").strip()
            if not filename:
                continue
            result.append({
                "filename": filename,
                "number": str(row.get("Номер ДБН") or "").strip(),
                "title": str(row.get("Найменування ДБН") or "").strip(),
                "category": str(row.get("Категорія") or "").strip(),
                "edessb_url": str(row.get("Посилання ЄДЕССБ") or "").strip(),
                "path": f"building_codes/{filename}",
            })
        workbook.close()
        return result
    return json.loads(path.read_text(encoding="utf-8"))


def open_document(path: Path) -> None:
    if os.name == "nt":
        os.startfile(path)  # type: ignore[attr-defined]
    elif os.name == "posix":
        subprocess.Popen(["open" if __import__('sys').platform == "darwin" else "xdg-open", str(path)])
