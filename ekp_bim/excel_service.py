from __future__ import annotations

from pathlib import Path
from typing import Iterable

HEADERS = ["guid", "type", "name", "id", "layer", "floor", "width", "height", "area", "note"]


def export_rows(path: Path, rows: Iterable[dict], sheet_name: str = "Elements") -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    ws.append(HEADERS)
    for row in rows:
        ws.append([row.get(key, "") for key in HEADERS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F6FEB")
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        width = min(42, max(12, max(len(str(c.value or "")) for c in column) + 2))
        ws.column_dimensions[column[0].column_letter].width = width
    wb.save(path)


def import_rows(path: Path) -> list[dict]:
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    values = list(ws.iter_rows(values_only=True))
    if not values:
        return []
    headers = [str(value or "").strip().lower() for value in values[0]]
    return [
        {headers[i]: value if value is not None else "" for i, value in enumerate(row) if i < len(headers)}
        for row in values[1:]
    ]


def demo_rows(kind: str) -> list[dict]:
    types = {"all": ["Wall", "Door", "Window"], "doors": ["Door"], "windows": ["Window"]}[kind]
    return [{"guid": f"DEMO-{i:04d}", "type": t, "name": {"Wall":"Стіна", "Door":"Двері", "Window":"Вікно"}[t],
             "id": f"{t[:1]}-{i:03d}", "layer": "ARCH", "floor": "1 поверх",
             "width": 900 if t == "Door" else 1500 if t == "Window" else "",
             "height": 2100 if t == "Door" else 1400 if t == "Window" else "", "area": "", "note": "Демо-дані"}
            for i, t in enumerate(types, 1)]
